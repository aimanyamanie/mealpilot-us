"""The review round trip: a filled-in workbook becomes sign-offs, and a
sign-off only holds while the thing it approved is unchanged."""
import csv
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import make_review_sheet  # noqa: E402
from apply_review import merge, read_workbook  # noqa: E402
from build_seed import apply_portion_signoff, apply_signoff  # noqa: E402

PIPELINE = Path(__file__).resolve().parents[1]


def seed():
    with (PIPELINE / "ingredients_seed.csv").open(encoding="utf-8", newline="") as f:
        return {r["slug"]: r for r in csv.DictReader(f)}


def filled_workbook(tmp_path, ingredient_rows, portion_rows=()):
    path = make_review_sheet.main(tmp_path)
    wb = load_workbook(path)
    ws = wb["Ingredients"]
    rows = {ws.cell(row=r, column=1).value: r for r in range(2, ws.max_row + 1)}
    for slug, values in ingredient_rows.items():
        for col, v in zip("VWXYZ", values[:5]):
            ws[f"{col}{rows[slug]}"] = v
        if len(values) > 5:
            ws[f"AA{rows[slug]}"] = values[5]
    ws = wb["Portions"]
    for r, status, reviewer, comment in portion_rows:
        ws[f"H{r}"], ws[f"I{r}"], ws[f"J{r}"] = status, reviewer, comment
    wb.save(path)
    return path


def test_approvals_and_fixes_are_read(tmp_path):
    path = filled_workbook(tmp_path, {
        "onion-yellow": ("Approve", "Approve", "Approve", "Aiman", "2026-10-30"),
        "worcestershire": ("Approve", "Fix", "Approve", "Aiman", "", "label has no wheat"),
    }, [(2, "Approve", "Aiman", "")])
    ingredients, portions, problems = read_workbook(path, seed())
    assert problems == []
    by_slug = {r["slug"]: r for r in ingredients}
    onion = by_slug["onion-yellow"]
    assert onion["fdc_id"] == "790646"
    assert onion["flags_snapshot"] == "none | plant"
    assert by_slug["worcestershire"]["allergens"] == "Fix"
    assert len(portions) == 1 and portions[0]["status"] == "Approve"


def test_bad_rows_are_refused(tmp_path):
    path = filled_workbook(tmp_path, {
        "garlic": ("Approve", "", "", "", ""),                     # no reviewer
        "carrot": ("Fix", "", "", "Aiman", ""),                     # fix without comment
        "celery": ("Yes", "", "", "Aiman", ""),                     # not a dropdown value
    })
    _, _, problems = read_workbook(path, seed())
    assert len(problems) == 3


def row_for(slug):
    """A seed row as build_seed holds it: booleans, not strings."""
    r = dict(seed()[slug])
    for k, v in r.items():
        if v in ("True", "False"):
            r[k] = v == "True"
    r["needs_verification"], r["allergens_reviewed"] = True, False
    return r


def signoff(**over):
    base = {"slug": "onion-yellow", "fdc_id": "790646", "nutrition": "Approve",
            "flags_snapshot": "none | plant", "allergens": "Approve", "diet": "Approve",
            "reviewer": "Aiman", "reviewed_on": "2026-10-30", "comment": ""}
    return {**base, **over}


def test_signoff_applies_to_what_was_approved():
    row = row_for("onion-yellow")
    assert apply_signoff(row, signoff(), 790646) == []
    assert row["needs_verification"] is False and row["verified_by"] == "Aiman"
    assert row["allergens_reviewed"] is True


def test_signoff_stops_applying_when_the_record_or_flags_change():
    row = row_for("onion-yellow")
    stale = apply_signoff(row, signoff(fdc_id="170000", flags_snapshot="milk | dairy"), 790646)
    assert len(stale) == 2
    assert row["needs_verification"] is True and row["allergens_reviewed"] is False


def test_allergens_need_diet_approved_too():
    row = row_for("onion-yellow")
    apply_signoff(row, signoff(diet="Fix"), 790646)
    assert row["allergens_reviewed"] is False


def test_portion_signoff_needs_same_grams():
    portion = {"slug": "garlic", "unit": "clove", "modifier": "", "grams": 3.0,
               "reviewed_by": "", "reviewed_at": ""}
    approved = {"status": "Approve", "grams": "3", "reviewer": "Aiman", "reviewed_on": "2026-10-30"}
    assert apply_portion_signoff(dict(portion), {**approved, "grams": "4"}) != []
    apply_portion_signoff(portion, approved)
    assert portion["reviewed_by"] == "Aiman"


def test_blank_decisions_keep_earlier_signoffs():
    earlier = signoff()
    later = signoff(nutrition="", allergens="", diet="", fdc_id="999", flags_snapshot="x")
    merged = merge(earlier, later)
    assert merged["nutrition"] == "Approve" and merged["fdc_id"] == "790646"
    assert merged["flags_snapshot"] == "none | plant"
